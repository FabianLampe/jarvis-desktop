"""Feste Datenoperationen, kein eval/Modellcode. Wird mit python -I gestartet."""
from __future__ import annotations

import json
import math
import csv
import sys
from pathlib import Path


def table(req):
    import pandas as pd
    p = Path(req["source"])
    if p.suffix.lower() == ".xlsx":
        df = pd.read_excel(p, sheet_name=req.get("sheet", 0), nrows=100001)
    else:
        with p.open(encoding="utf-8-sig") as f:
            sample = f.read(8192)
        try:
            sep = "\t" if p.suffix.lower() == ".tsv" else csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            sep = ","  # auch einspaltige Dateien korrekt lesen
        df = pd.read_csv(p, sep=sep, encoding="utf-8-sig", nrows=100001)
    if len(df) > 100000 or len(df.columns) > 100:
        raise ValueError("Maximal 100000 Zeilen und 100 Spalten unterstützt; Datei vorher aufteilen")
    return df


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def document(req):
    title, paragraphs, rows = str(req["title"]), req["paragraphs"], req["rows"]
    if not isinstance(paragraphs, list) or not all(isinstance(p, str) for p in paragraphs):
        raise ValueError("paragraphs muss eine Liste von Textabsätzen sein")
    if (not isinstance(rows, list) or len(rows) > 5000 or
            any(not isinstance(r, list) or not 1 <= len(r) <= 100 for r in rows)):
        raise ValueError("Tabelle maximal 5000 Zeilen und 100 Spalten")
    if sum(len(r) for r in rows) > 20000:
        raise ValueError("Tabelle maximal 20000 Zellen; Daten vorher zusammenfassen")
    if sum(len(p) for p in paragraphs) > 50000:
        raise ValueError("Text maximal 50000 Zeichen")
    p = Path(req["output"])
    if p.suffix.lower() == ".docx":
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for text in paragraphs:
            doc.add_paragraph(text)
        if rows:
            t = doc.add_table(rows=0, cols=max(len(r) for r in rows))
            t.style = "Light Shading Accent 1"
            for row in rows:
                cells = t.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = str(value if value is not None else "")
        doc.save(p)
    elif p.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Auswertung"
        for row in [[title], *[[x] for x in paragraphs], [], *rows]:
            # Text aus Datenquellen darf keine Excel-Formel werden.
            ws.append([("'" + v) if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v for v in row])
        ws.cell(1, 1).font = Font(bold=True, size=16)
        ws.cell(1, 1).fill = PatternFill("solid", fgColor="DDEBF7")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(14, max(len(str(c.value or "")) for c in col) + 2))
        wb.save(p)
    elif p.suffix.lower() == ".pdf":
        from html import escape
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import matplotlib
        fonts = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        pdfmetrics.registerFont(TTFont("JarvisSans", str(fonts / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("JarvisSansBold", str(fonts / "DejaVuSans-Bold.ttf")))
        styles = getSampleStyleSheet()
        styles["Title"].fontName = "JarvisSansBold"
        styles["BodyText"].fontName = "JarvisSans"
        story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
        story += [Paragraph(escape(t).replace("\n", "<br/>"), styles["BodyText"]) for t in paragraphs]
        if rows:
            cols = max(len(r) for r in rows)
            if cols > 8:
                raise ValueError("PDF-Tabellen maximal 8 Spalten; für breite Tabellen XLSX verwenden")
            values = [[Paragraph(escape(str(v if v is not None else "")), styles["BodyText"]) for v in r]
                      + [""] * (cols - len(r)) for r in rows]
            t = Table(values, colWidths=[(A4[0] - 80) / cols] * cols, repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEBF7")),
                                   ("GRID", (0, 0), (-1, -1), .3, colors.lightgrey),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.extend([Spacer(1, 12), t])
        SimpleDocTemplate(str(p), pagesize=A4, leftMargin=40, rightMargin=40).build(story)
    else:
        raise ValueError("Format nicht unterstützt")
    return {"format": p.suffix[1:]}


def main(req):
    if req["action"] == "document":
        return document(req)
    df = table(req)
    if req["action"] == "analyze":
        if req.get("columns"):
            columns = req["columns"]
            if not isinstance(columns, list) or any(c not in df.columns for c in columns):
                raise ValueError("Unbekannte Spaltenauswahl")
            df = df[columns]
        numeric = df.select_dtypes(include="number")
        return clean({"rows": len(df), "columns": [str(c) for c in df.columns],
                      "dtypes": df.dtypes.astype(str).to_dict(), "missing": df.isna().sum().to_dict(),
                      "numeric": numeric.describe().to_dict() if len(numeric.columns) and len(df) else {}})
    if req["action"] != "plot" or req.get("kind") not in ("line", "bar", "hist"):
        raise ValueError("Unbekannte Datenoperation")
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    y, x = req["y"], req.get("x")
    if y not in df or (x and x not in df):
        raise ValueError("Unbekannte Spalte; zuerst analyze_table verwenden")
    if not __import__("pandas").api.types.is_numeric_dtype(df[y]):
        raise ValueError("y muss eine numerische Spalte sein")
    if req["kind"] == "bar" and len(df) > 100:
        raise ValueError("Balkendiagramme maximal 100 Zeilen; Daten vorher aggregieren")
    if req["kind"] == "line" and len(df) > 10000:
        raise ValueError("Liniendiagramme maximal 10000 Zeilen; Daten vorher aggregieren")
    fig, ax = plt.subplots(figsize=(10, 5.5), layout="constrained")
    if req["kind"] == "hist":
        ax.hist(df[y].dropna(), bins=30)
    elif req["kind"] == "bar":
        ax.bar(df[x].astype(str) if x else df.index, df[y])
    else:
        ax.plot(df[x] if x else df.index, df[y])
    ax.set(title=req["title"], xlabel=x or "Zeile", ylabel=y)
    ax.grid(alpha=.2)
    fig.savefig(req["output"], dpi=140)
    plt.close(fig)
    return {"rows": len(df), "kind": req["kind"]}


if __name__ == "__main__":
    try:
        data = main(json.loads(Path(sys.argv[1]).read_text(encoding="utf-8")))
        print(json.dumps({"ok": True, **data}, ensure_ascii=False, allow_nan=False, default=str))
    except ImportError as exc:
        print(json.dumps({"ok": False, "error": f"Abhängigkeit fehlt ({exc.name}); pip install -r requirements-tools.txt"}))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)[:1000]}))
