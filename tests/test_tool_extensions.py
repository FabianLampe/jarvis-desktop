"""Neue Werkzeuge gegen echte Git-Repos, pytest, Dateien und einen stdio-MCP-Server."""
import json
import subprocess
import sys
import threading
import time
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pytest

from jarvis_core import config
from jarvis_core.tools import build_registry
from jarvis_core.tools import artifact_tools, mac_ui_tools, mcp_tools, remote_tools, workflow_tools
from jarvis_core.tools.registry import ToolError
from tests.tools_helpers import run, runtime


def call(tool_name, **args):
    return {"function": {"name": tool_name, "arguments": args}}


def execute(rt, tool_name, **args):
    return rt.registry.execute(rt.ctx, tool_name, args, 20000)


def configure(rt, **extensions):
    cfg = replace(rt.cfg, extensions=extensions)
    rt.ctx.cfg = rt.agent.cfg = cfg
    return cfg


def test_browser_snapshot_is_reachable_and_untrusted(tmp_path):
    rt = runtime(tmp_path)
    class Browser:
        def ensure_started(self): pass
        def browser_snapshot(self): return 'uid=2 button "Speichern"'
    rt.ctx.services["browser"] = Browser()
    res = execute(rt, "browser_snapshot")
    assert res.ok and "uid=2" in res.text and rt.ctx.tainted


def test_git_tools_use_a_real_repository(tmp_path):
    rt = runtime(tmp_path)
    repo = rt.ctx.workspace / "project"
    repo.mkdir()
    subprocess.run(["git", "init", "-q", str(repo)], check=True)
    for key, value in [("user.name", "Test"), ("user.email", "test@example.invalid")]:
        subprocess.run(["git", "-C", str(repo), "config", key, value], check=True)
    (repo / "a.txt").write_text("old\n")
    subprocess.run(["git", "-C", str(repo), "add", "a.txt"], check=True)
    subprocess.run(["git", "-C", str(repo), "commit", "-qm", "initial"], check=True)
    (repo / "a.txt").write_text("new\n")
    assert "a.txt" in execute(rt, "git_status", repo="project").data["stdout"]
    diff = execute(rt, "git_diff", repo="project")
    assert diff.ok and "+new" in diff.data["stdout"] and "-old" in diff.data["stdout"]
    assert rt.registry.assess(rt.ctx, "git_branch", {"name": "feature/test"}).level == "confirm"
    assert execute(rt, "git_branch", repo="project", name="feature/test").ok
    assert "feature/test" in execute(rt, "git_status", repo="project").data["stdout"]
    assert not execute(rt, "git_branch", repo="project", name="--force").ok
    assert not execute(rt, "git_diff", repo="project", path="../../escape").ok


def test_real_pytest_counts_failures_and_rejects_option_injection(tmp_path):
    rt = runtime(tmp_path)
    (rt.ctx.workspace / "test_sample.py").write_text("def test_ok(): assert 1 == 1\ndef test_bad(): assert 1 == 2\n")
    assert rt.registry.assess(rt.ctx, "run_tests", {}).level == "confirm"
    res = execute(rt, "run_tests", targets=["test_sample.py"])
    assert not res.ok and res.data["tests"] == 2 and res.data["failures"] == 1
    good = execute(rt, "run_tests", targets=["test_sample.py::test_ok"])
    assert good.ok and good.data["tests"] == 1
    assert not execute(rt, "run_tests", targets=["--override-ini=x"]).ok
    assert not execute(rt, "run_tests", targets=["../outside.py"]).ok


def test_file_verification_checks_real_content_and_protected_paths(tmp_path):
    rt = runtime(tmp_path)
    (rt.ctx.workspace / "a.txt").write_text("Hallo")
    assert execute(rt, "verify_file", path="a.txt", contains="Hallo", min_bytes=3).ok
    assert not execute(rt, "verify_file", path="a.txt", contains="Nein").ok
    assert not execute(rt, "verify_file", path="a.txt", sha256="wrong").ok
    assert not execute(rt, "verify_file", path="missing").ok
    assert not execute(rt, "verify_file", path=str(rt.ctx.data_dir / "secret")).ok


def test_github_write_policy_and_structured_payload(tmp_path, monkeypatch):
    rt = runtime(tmp_path)
    monkeypatch.setenv("JARVIS_GITHUB_TOKEN", "secret-test-token")
    seen = []
    def request(method, url, **kw):
        seen.append((method, url, kw))
        return SimpleNamespace(ok=True, json=lambda: {"number": 4, "title": "Fix", "html_url": "https://example.invalid/pr/4"})
    monkeypatch.setattr(remote_tools.requests, "request", request)
    a = {"action": "create_pr", "repo": "owner/repo", "title": "Fix", "head": "feature", "base": "main"}
    assert rt.registry.assess(rt.ctx, "github", a).level == "confirm"
    res = execute(rt, "github", **a)
    assert res.ok and seen[0][0] == "POST" and seen[0][2]["json"]["draft"] is True
    assert "secret-test-token" not in res.text
    assert rt.registry.assess(rt.ctx, "github", {"action": "issues"}).level == "allow"
    assert not execute(rt, "github", repo="https://evil.invalid", action="issues").ok
    monkeypatch.delenv("JARVIS_GITHUB_TOKEN")
    assert execute(rt, "github", action="issues").ok  # öffentliche Repositories auch ohne Token
    assert not execute(rt, "github", **a).ok


def test_ssh_requires_configured_alias_and_uses_strict_keys(tmp_path, monkeypatch):
    rt = runtime(tmp_path)
    assert not execute(rt, "ssh_diagnose", host="pi").ok
    configure(rt, ssh={"enabled": True, "hosts": {"pi": {"hostname": "pi.local", "user": "tester", "port": 2222}}})
    seen = []
    def invoke(ctx, argv, **kw):
        seen.append(argv)
        return SimpleNamespace(ok=True, exit_code=0, stdout="healthy", stderr="", timed_out=False, cancelled=False, truncated=False, error=None)
    monkeypatch.setattr(remote_tools, "run", invoke)
    assert execute(rt, "ssh_diagnose", host="pi", check="network").ok
    assert "StrictHostKeyChecking=yes" in seen[0] and "BatchMode=yes" in seen[0]
    assert seen[0][-2] == "tester@pi.local" and "status=0" in seen[0][-1]
    assert not execute(rt, "ssh_diagnose", host="pi", check="logs", service="x; reboot").ok
    assert not execute(rt, "ssh_diagnose", host="unknown").ok
    assert rt.registry.assess(rt.ctx, "ssh_run", {"host": "pi", "argv": ["echo", "x"]}).level == "confirm"
    assert execute(rt, "ssh_run", host="pi", argv=["echo", "$(touch nope); x"]).ok
    assert seen[-1][-1] == "echo '$(touch nope); x'"


def test_csv_analysis_and_chart_are_real_files(tmp_path):
    rt = runtime(tmp_path)
    (rt.ctx.workspace / "a.csv").write_text("tag,wert\nA,1\nB,3\nC,\n")
    res = execute(rt, "analyze_table", path="a.csv")
    assert res.ok, res.text
    assert res.data["rows"] == 3 and res.data["missing"]["wert"] == 1
    assert res.data["numeric"]["wert"]["mean"] == 2
    plot = execute(rt, "plot_table", path="a.csv", output="plot.png", x="tag", y="wert", kind="bar")
    assert plot.ok, plot.text
    assert (rt.ctx.workspace / "plot.png").read_bytes().startswith(b"\x89PNG")
    assert not execute(rt, "plot_table", path="a.csv", output="bad.png", y="missing").ok
    assert not (rt.ctx.workspace / "bad.png").exists()
    assert not list(rt.ctx.workspace.glob(".jarvis-*"))


def test_single_column_csv_and_selected_columns(tmp_path):
    rt = runtime(tmp_path)
    (rt.ctx.workspace / "one.csv").write_text("wert\n1\n2\n3\n")
    res = execute(rt, "analyze_table", path="one.csv", columns=["wert"])
    assert res.ok and res.data["columns"] == ["wert"] and res.data["numeric"]["wert"]["mean"] == 2
    assert not execute(rt, "analyze_table", path="one.csv", columns=["missing"]).ok


@pytest.mark.parametrize("suffix", ["pdf", "docx", "xlsx"])
def test_document_export_and_overwrite_policy(tmp_path, suffix):
    rt = runtime(tmp_path)
    res = execute(rt, "create_document", output=f"report.{suffix}", title="Prüfbericht",
                  paragraphs=["Ein Absatz mit Umlauten: äöü."], rows=[["Spalte", "Wert"], ["A", 42]])
    assert res.ok, res.text
    p = rt.ctx.workspace / f"report.{suffix}"
    assert p.stat().st_size > 1000
    if suffix == "pdf":
        assert p.read_bytes().startswith(b"%PDF")
    elif suffix == "docx":
        from docx import Document
        assert "Prüfbericht" in Document(p).paragraphs[0].text
    else:
        from openpyxl import load_workbook
        assert load_workbook(p).active["A1"].value == "Prüfbericht"
    assert not execute(rt, "create_document", output=p.name).ok
    assert rt.registry.assess(rt.ctx, "create_document", {"output": p.name, "overwrite": True}).level == "confirm"


def test_excel_output_keeps_formula_like_text_literal(tmp_path):
    from openpyxl import load_workbook
    rt = runtime(tmp_path)
    res = execute(rt, "create_document", output="a.xlsx", rows=[["=1+1", "@x", 3]])
    assert res.ok, res.text
    ws = load_workbook(rt.ctx.workspace / "a.xlsx").active
    assert all(c.data_type != "f" for row in ws for c in row)


def test_cancelled_document_does_not_replace_existing_file(tmp_path, monkeypatch):
    rt = runtime(tmp_path)
    p = rt.ctx.workspace / "old.pdf"
    p.write_bytes(b"original")
    def worker(ctx, req):
        Path(req["output"]).write_bytes(b"new")
        ctx.cancel.set()
        return {}
    monkeypatch.setattr(artifact_tools, "worker", worker)
    assert not execute(rt, "create_document", output="old.pdf", overwrite=True).ok
    assert p.read_bytes() == b"original" and not list(p.parent.glob(".jarvis-*"))


def test_workflow_persists_and_executes_each_step(tmp_path):
    rt = runtime(tmp_path, {"tool_calls": [call("workflow_run", name="project")]}, {"content": "Datei geprüft."})
    steps = [{"tool": "make_dir", "args": {"path": "project"}},
             {"tool": "write_file", "args": {"path": "project/a.txt", "content": "Hallo"}},
             {"tool": "verify_file", "args": {"path": "project/a.txt", "contains": "Hallo"}}]
    assert execute(rt, "workflow_save", name="project", steps=steps).ok
    ev = run(rt, "Starte den Workflow project")
    assert ev[-1]["ok"] and (rt.ctx.workspace / "project/a.txt").read_text() == "Hallo"
    assert [e["tool"] for e in ev if e["type"] == "result"] == ["make_dir", "write_file", "verify_file"]
    assert execute(rt, "workflow_list").data["workflows"] == ["project"]


def test_workflow_cannot_bypass_confirmation_or_recurse(tmp_path):
    rt = runtime(tmp_path, {"tool_calls": [call("workflow_run", name="delete")]})
    (rt.ctx.workspace / "a.txt").write_text("keep")
    steps = [{"tool": "delete_path", "args": {"path": "a.txt"}}, {"tool": "make_dir", "args": {"path": "never"}}]
    assert execute(rt, "workflow_save", name="delete", steps=steps).ok
    ev = run(rt, "Starte den Workflow delete", answer=lambda _: False)
    assert not ev[-1]["ok"] and (rt.ctx.workspace / "a.txt").exists()
    assert not (rt.ctx.workspace / "never").exists()
    assert not execute(rt, "workflow_save", name="bad", steps=[{"tool": "workflow_run", "args": {"name": "bad"}}]).ok
    assert not execute(rt, "workflow_save", name="../escape", steps=steps).ok
    assert not execute(rt, "workflow_save", name="bad", steps=[{"tool": "not-real", "args": {}}]).ok


def test_workflow_stops_at_failure_and_obeys_shared_step_budget(tmp_path):
    rt = runtime(tmp_path, {"tool_calls": [call("workflow_run", name="fail")]})
    steps = [{"tool": "verify_file", "args": {"path": "missing"}}, {"tool": "make_dir", "args": {"path": "never"}}]
    assert execute(rt, "workflow_save", name="fail", steps=steps).ok
    assert not run(rt, "Workflow fail ausführen")[-1]["ok"]
    assert not (rt.ctx.workspace / "never").exists()
    # Ein vormals gültiges längeres Rezept darf ein kleineres aktuelles Budget nicht umgehen.
    rt.agent.cfg = rt.ctx.cfg = replace(rt.cfg, tools=replace(rt.cfg.tools, max_steps=1))
    assert not execute(rt, "workflow_get", name="fail").ok


def test_native_ui_is_optional_and_actions_require_confirmation(tmp_path):
    rt = runtime(tmp_path)
    assert not execute(rt, "mac_ui_snapshot").ok
    configure(rt, mac_ui={"enabled": True})
    class FakeUI:
        def snapshot(self, app): return {"app": app, "elements": [{"id": "ax1", "role": "AXButton"}]}
        def action(self, a): return {"performed": a["action"]}
    rt.ctx.services["mac_ui"] = FakeUI()
    assert execute(rt, "mac_ui_snapshot", app="Finder").ok and rt.ctx.tainted
    assert rt.registry.assess(rt.ctx, "mac_ui_action", {"action": "click", "id": "ax1"}).level == "confirm"
    assert execute(rt, "mac_ui_action", action="click", id="ax1").ok


def test_native_controller_rejects_stale_ids_and_wrong_foreground():
    ui = mac_ui_tools.MacUI.__new__(mac_ui_tools.MacUI)
    ui.app = SimpleNamespace(processIdentifier=lambda: 1)
    ui.at = 0
    ui.elements = {}
    with pytest.raises(ToolError, match="frischen"):
        ui.action({"action": "click", "id": "old"})
    ui.at = time.monotonic()
    ui.apps = SimpleNamespace(NSWorkspace=SimpleNamespace(sharedWorkspace=lambda:
        SimpleNamespace(frontmostApplication=lambda: SimpleNamespace(processIdentifier=lambda: 2))))
    with pytest.raises(ToolError, match="gewechselt"):
        ui.action({"action": "key", "key": "cmd+s"})


def test_native_controller_snapshots_and_actions_use_ax_elements():
    app = SimpleNamespace(localizedName=lambda: "TestApp", processIdentifier=lambda: 123)
    workspace = SimpleNamespace(frontmostApplication=lambda: app, runningApplications=lambda: [app])
    tree = {1: {"AXRole": "AXApplication", "AXChildren": [2, 3]},
            2: {"AXRole": "AXButton", "AXTitle": "Save"},
            3: {"AXRole": "AXTextField", "AXSubrole": "AXSecureTextField", "AXValue": "secret"}}
    pressed = []
    ui = mac_ui_tools.MacUI.__new__(mac_ui_tools.MacUI)
    ui.ax = SimpleNamespace(AXUIElementCreateApplication=lambda pid: 1,
                           AXUIElementCopyAttributeValue=lambda element, key, _: (0, tree[element][key]) if key in tree[element] else (1, None),
                           AXUIElementPerformAction=lambda element, action: pressed.append((element, action)) or 0)
    ui.apps = SimpleNamespace(NSWorkspace=SimpleNamespace(sharedWorkspace=lambda: workspace))
    ui.elements, ui.app, ui.at = {}, None, 0
    snapshot = ui.snapshot("TestApp")
    assert "secret" not in json.dumps(snapshot)
    button = next(e for e in snapshot["elements"] if e["role"] == "AXButton")
    ui.action({"action": "click", "id": button["id"]})
    assert pressed == [(2, "AXPress")] and not ui.elements
    new = ui.snapshot("TestApp")
    assert button["id"] not in {e["id"] for e in new["elements"]}


@pytest.fixture
def mcp_runtime(tmp_path):
    rt = runtime(tmp_path)
    server = tmp_path / "mcp_server.py"
    server.write_text('''from mcp.server.fastmcp import FastMCP
import asyncio
m = FastMCP("test")
@m.tool()
def echo(text: str) -> str:
    return "echo: " + text
@m.tool()
def fail() -> str:
    raise ValueError("expected failure")
@m.tool()
async def slow() -> str:
    await asyncio.sleep(30)
    return "late"
if __name__ == "__main__": m.run(transport="stdio")
''')
    configure(rt, mcp={"enabled": True, "servers": {"test": {"command": sys.executable, "args": [str(server)]}}})
    return rt


def test_mcp_real_stdio_discovery_and_execution(mcp_runtime):
    rt = mcp_runtime
    listing = execute(rt, "mcp_list_tools", server="test")
    assert listing.ok, listing.text
    assert "echo" in {t["name"] for t in listing.data["tools"]}
    res = execute(rt, "mcp_call", server="test", tool="echo", arguments={"text": "Hallo"})
    assert res.ok and "echo: Hallo" in res.text
    assert rt.registry.assess(rt.ctx, "mcp_call", {"server": "test", "tool": "echo"}).level == "confirm"
    assert not execute(rt, "mcp_call", server="test", tool="echo", arguments={"text": 123}).ok
    assert not execute(rt, "mcp_call", server="test", tool="fail").ok
    assert not execute(rt, "mcp_call", server="unknown", tool="echo").ok


def test_mcp_cancel_stops_waiting_for_server(mcp_runtime):
    rt = mcp_runtime
    timer = threading.Timer(1, rt.ctx.cancel.set)
    started = time.monotonic()
    timer.start()
    try:
        res = execute(rt, "mcp_call", server="test", tool="slow")
    finally:
        timer.cancel()
    assert not res.ok and time.monotonic() - started < 8


@pytest.mark.parametrize("text,group", [("Pi Netzwerk diagnostizieren", "ssh"), ("GitHub Issues lesen", "git"),
    ("pytest ausführen", "testing"), ("CSV analysieren", "analysis"), ("Word Bericht erstellen", "documents_create"),
    ("Workflow speichern", "workflows"), ("MCP Werkzeuge anzeigen", "mcp"), ("Finder Fenster bedienen", "mac_ui")])
def test_new_groups_are_routed(text, group):
    from jarvis_core.tool_intent import select_groups
    assert group in select_groups(text)


def test_extension_config_roundtrip_and_validation(tmp_path):
    p = tmp_path / "c.yaml"
    p.write_text("extensions:\n  ssh:\n    enabled: true\n    hosts: {}\n")
    assert config.load(p).extensions["ssh"]["enabled"] is True
    p.write_text("extensions:\n  mac_ui:\n    enabled: yesplease\n")
    with pytest.raises(config.ConfigError):
        config.load(p)


def test_all_tool_schemas_validate():
    from jsonschema import Draft202012Validator
    for tool in build_registry().specs():
        Draft202012Validator.check_schema(tool["function"]["parameters"])


def test_large_structured_results_remain_valid_json(tmp_path):
    from jarvis_core.tools.extension_utils import result
    from jarvis_core.tools.registry import Tool
    rt = runtime(tmp_path)
    rt.registry.register(Tool("big", "test", "test", {}, (),
                              lambda c, a: result({"exit_code": 0, "stdout": '"x\\n' * 10000})))
    res = rt.registry.execute(rt.ctx, "big", {}, 600)
    data = json.loads(res.text)
    assert len(res.text) <= 600 and data["truncated"] and data["exit_code"] == 0


def test_setup_status_reports_configuration_without_connecting(tmp_path):
    from jarvis_core.tools.extension_utils import status_lines
    rt = runtime(tmp_path)
    cfg = configure(rt, ssh={"enabled": True, "hosts": {"pi": {"hostname": "pi.local", "user": "test"}}})
    text = "\n".join(status_lines(cfg))
    assert "Verbindung nicht geprüft" in text and "pi" in text and "MCP: deaktiviert" in text
